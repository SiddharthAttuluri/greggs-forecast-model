"""Monte Carlo simulation of the Greggs three-statement forecast (FY26–FY28).

How it works
------------
1. `project()` re-implements the Excel model's FY26–FY28 logic in Python: the revenue
   build (shop roll-forward, existing estate, annualisation, new shops, B2B), the income
   statement, the supporting schedules (fixed assets, leases, restructuring provision,
   deferred tax, dividends), the balance sheet and the cash flow statement.
2. Control: run with the base-case drivers, it must reproduce the Excel model's base
   case (checked in `check_against_excel()`), and the balance sheet must balance in
   every year of every run.
3. Each uncertain driver gets a triangular distribution: min / most likely (base) / max,
   read from the Drivers tab. One draw per driver per run is applied to all three years
   (a driver that runs hot in FY26 stays hot), scaled to each year's own range.
4. Links between drivers (like-for-like growth with margin, and with B2B growth) are
   imposed with a Gaussian copula: correlated normal draws are converted to uniforms,
   then to triangular values.
5. 10,000 runs; results summarised as percentiles, probabilities and a tornado chart
   (one driver at a time moved to its min and max, others at base).

Usage:
    python simulation/monte_carlo.py --workbook model/GRG_model_v0.7.xlsx
"""
import argparse
import csv
import pathlib

import numpy as np
from openpyxl import load_workbook
from scipy.stats import norm

ROOT = pathlib.Path(__file__).resolve().parents[1]
YEARS = ["FY26", "FY27", "FY28"]
RNG_SEED = 42
N_RUNS = 10_000

# drivers simulated: name -> Drivers tab row
SIM_DRIVERS = {"open": 6, "cmshare": 7, "lfl": 8, "newshop": 9, "b2b": 10, "margin": 11,
               "capex": 15, "tax": 18}
LABELS = {"open": "Net shop openings", "cmshare": "Company-managed share of openings",
          "lfl": "Like-for-like growth", "newshop": "New-shop sales vs mature shop",
          "b2b": "B2B growth", "margin": "Underlying operating margin", "capex": "Capex",
          "tax": "Tax rate", "rtotal": "Restructuring total cost"}


# ---------------------------------------------------------------- inputs
def load_inputs(workbook):
    wb = load_workbook(workbook, data_only=True)
    d = wb["Drivers"]

    def row(r):
        return [d.cell(row=r, column=c).value for c in (3, 4, 5)]

    def rng(r):
        v = [d.cell(row=r, column=c).value for c in range(6, 12)]
        return [(v[0], v[1]), (v[2], v[3]), (v[4], v[5])]

    base = {k: row(r) for k, r in SIM_DRIVERS.items()}
    ranges = {k: rng(r) for k, r in SIM_DRIVERS.items()}
    base["rtotal"] = [d["C12"].value] * 3
    ranges["rtotal"] = [(d["F12"].value, d["G12"].value)] * 3
    fixed = {
        "rpl": row(13), "rcash": row(14), "dps": row(17), "gm": row(21), "admpct": row(22),
        "invd": row(23), "recd": row(24), "payd": row(25), "deppct": row(26), "intshare": row(27),
        "amortpct": row(28), "roudep": row(29), "leaseint": row(30), "cashint": row(31),
        "borrint": row(32), "dtlpct": row(33), "shares": row(34), "curlease": row(35),
        "firstyr": row(36),
    }
    corr = {("lfl", "margin"): d["C38"].value, ("lfl", "b2b"): d["C39"].value}

    hist = {}
    with open(ROOT / "data" / "greggs_historicals_fy21_fy25.csv", newline="") as f:
        for r in csv.DictReader(f):
            hist[(r["key"], r["year"])] = float(r["value"])
    return base, ranges, fixed, corr, hist


def opening_position(hist):
    g = lambda k, y="FY25": hist.get((k, y), 0.0)  # noqa: E731
    cm25 = g("shops") - g("fr")
    cm24 = g("shops", "FY24") - g("fr", "FY24")
    return {
        "cm": cm25, "cmnet": cm25 - cm24, "cmrev": g("rev_cm"), "b2b": g("rev_b2b"),
        "sps": g("rev_cm") / ((cm24 + cm25) / 2), "shops": g("shops"), "fr": g("fr"),
        "intang": g("intang"), "ppe": g("ppe"), "rou": g("rou"), "inv": g("inv"), "rec": g("rec"),
        "cash": g("cash"), "pay": g("pay"), "ctl": g("ctl"), "lease": -(g("leasec") + g("leasenc")),
        "provc": g("provc"), "borr": g("borr"), "dtl": g("dtl"), "re": g("re"), "dps": 0.69,
        "other_liab": g("opay") + g("provnc") + g("penl"), "equity_other": g("sc") + g("sp") + g("crr"),
    }


# ---------------------------------------------------------------- the model
def project(dv, fx, op):
    """Run FY26–FY28. dv: simulated drivers (lists of 3), fx: fixed drivers, op: FY25 position."""
    s = dict(op)
    out = {k: [] for k in ("rev", "upbt", "pbt", "pat", "fcf", "cash", "netcash", "bs_diff", "uop")}
    for j in range(3):
        f = {k: v[j] for k, v in fx.items()}
        x = {k: v[j] for k, v in dv.items()}
        # revenue build (Revenue tab)
        cmnet = x["open"] * x["cmshare"]
        sps = s["sps"] * (1 + x["lfl"])
        cmrev = s["cmrev"] * (1 + x["lfl"]) + (s["cmnet"] * (1 - f["firstyr"]) + cmnet * f["firstyr"]) * sps * x["newshop"]
        b2b = s["b2b"] * (1 + x["b2b"])
        rev = cmrev + b2b
        shops = s["shops"] + x["open"]
        fr = s["fr"] + (1 - x["cmshare"]) * x["open"]
        cm = shops - fr
        # income statement
        cos = -rev * (1 - f["gm"])
        rpl = x["rtotal"] * f["rpl"]
        rcash = x["rtotal"] * f["rcash"]
        uop = x["margin"] * rev
        op_ = uop - rpl
        cashint = np.maximum(s["cash"], 0) * f["cashint"]
        leaseint = s["lease"] * f["leaseint"]
        borrint = -s["borr"] * f["borrint"]
        pbt = op_ + cashint - leaseint - borrint
        tax = -x["tax"] * pbt
        pat = pbt + tax
        # schedules
        capint = x["capex"] * f["intshare"]
        capppe = x["capex"] - capint
        amort = s["intang"] * f["amortpct"]
        dep = s["ppe"] * f["deppct"]
        lease = s["lease"] * cm / s["cm"]
        roudep = s["rou"] * f["roudep"]
        rou = s["rou"] * cm / s["cm"]
        princ = s["lease"] + (rou - s["rou"] + roudep) - lease
        dtl = s["dtl"] - f["dtlpct"] * (x["capex"] - dep - amort)
        div = s["dps"] * f["shares"]
        # balance sheet items
        intang = s["intang"] + capint - amort
        ppe = s["ppe"] + capppe - dep
        inv = -cos * f["invd"] / 365
        rec = rev * f["recd"] / 365
        pay = -rev * f["payd"] / 365
        provc = s["provc"] - (rpl - rcash)
        re = s["re"] + pat - div
        # cash flow
        cgo = (op_ + dep + amort + roudep - (inv - s["inv"]) - (rec - s["rec"])
               - (pay - s["pay"]) - (provc - s["provc"]))
        taxp = tax - (dtl - s["dtl"])
        cfo = cgo + taxp + cashint - leaseint - borrint
        cash = s["cash"] + cfo - x["capex"] - div - princ
        # balance check: assets - liabilities - equity
        assets = intang + ppe + rou + inv + rec + cash
        liabs = pay + s["ctl"] - lease + provc + s["borr"] + dtl + s["other_liab"]
        bs_diff = assets + liabs - (s["equity_other"] + re)
        for k, v in (("rev", rev), ("upbt", pbt + rpl), ("pbt", pbt), ("pat", pat), ("uop", uop),
                     ("fcf", cfo - x["capex"] - princ), ("cash", cash), ("netcash", cash + s["borr"]),
                     ("bs_diff", bs_diff)):
            out[k].append(v)
        s.update(cm=cm, cmnet=cmnet, cmrev=cmrev, b2b=b2b, sps=sps, shops=shops, fr=fr, intang=intang,
                 ppe=ppe, rou=rou, inv=inv, rec=rec, cash=cash, pay=pay, lease=lease, provc=provc,
                 dtl=dtl, re=re, dps=f["dps"])
    return out


# ---------------------------------------------------------------- sampling
def tri_ppf(u, lo, mode, hi):
    """Inverse CDF of the triangular distribution (vectorised)."""
    lo, mode, hi = float(lo), float(mode), float(hi)
    if hi == lo:
        return np.full_like(u, mode)
    c = (mode - lo) / (hi - lo)
    return np.where(u < c, lo + np.sqrt(u * (hi - lo) * (mode - lo)),
                    hi - np.sqrt((1 - u) * (hi - lo) * (hi - mode)))


def sample(base, ranges, corr, n, seed=RNG_SEED):
    names = list(base.keys())
    k = len(names)
    c = np.eye(k)
    for (a, b), rho in corr.items():
        i, j = names.index(a), names.index(b)
        c[i, j] = c[j, i] = rho
    z = np.random.default_rng(seed).standard_normal((n, k)) @ np.linalg.cholesky(c).T
    u = norm.cdf(z)
    draws = {}
    for i, name in enumerate(names):
        draws[name] = [tri_ppf(u[:, i], ranges[name][j][0], base[name][j], ranges[name][j][1]) for j in range(3)]
    return draws


# ---------------------------------------------------------------- control
def check_against_excel(workbook, base_out):
    wb = load_workbook(workbook, data_only=True)
    fs = wb["Forecast"]
    rows = {}
    for r in range(1, fs.max_row + 1):
        lab = fs.cell(row=r, column=1).value
        if lab and lab not in rows:
            rows[lab] = r
    pairs = [("Revenue", "rev"), ("Underlying profit before tax", "upbt"), ("Profit before tax", "pbt"),
             ("Cash and cash equivalents", "cash"), ("Free cash flow after leases", "fcf")]
    worst = 0.0
    table = []
    for lab, key in pairs:
        excel = [fs.cell(row=rows[lab], column=c).value for c in (8, 9, 10)]
        diff = [p - e for p, e in zip(base_out[key], excel)]
        worst = max(worst, max(abs(d) for d in diff))
        table.append((lab, excel, base_out[key], diff))
    return worst, table


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default=str(ROOT / "model" / "GRG_model_v0.7.xlsx"))
    ap.add_argument("--runs", type=int, default=N_RUNS)
    ap.add_argument("--out", default=str(ROOT / "simulation" / "results"))
    args = ap.parse_args()
    outdir = pathlib.Path(args.out)
    outdir.mkdir(parents=True, exist_ok=True)

    base, ranges, fixed, corr, hist = load_inputs(args.workbook)
    op = opening_position(hist)

    base_out = project(base, fixed, op)
    worst, table = check_against_excel(args.workbook, base_out)
    print("Control: Python base case vs Excel base case (FY26, FY27, FY28)")
    for lab, e, p, d in table:
        print(f"  {lab:<32} Excel {[round(v, 1) for v in e]}  Python {[round(float(v), 1) for v in p]}")
    print(f"  largest difference: £{worst:.4f}m  ->  {'PASS' if worst < 0.05 else 'FAIL'}")

    draws = sample(base, ranges, corr, args.runs)
    sims = project(draws, fixed, op)
    bs_worst = max(float(np.max(np.abs(a))) for a in sims["bs_diff"])
    print(f"Balance check across {args.runs:,} runs x 3 years: worst difference £{bs_worst:.6f}m")

    pct = [5, 25, 50, 75, 95]
    summary = []
    for key, lab in (("rev", "Revenue"), ("upbt", "Underlying profit before tax"),
                     ("fcf", "Free cash flow after leases"), ("cash", "Closing cash")):
        for j, y in enumerate(YEARS):
            a = sims[key][j]
            summary.append([lab, y] + [round(float(np.percentile(a, p)), 1) for p in pct]
                           + [round(float(np.mean(a)), 1), round(base_out[key][j], 1)])
    min_cash = np.minimum.reduce(sims["cash"])
    probs = [
        ("FY26 underlying PBT above FY25 (£171.9m)", float(np.mean(sims["upbt"][0] > 171.9))),
        ("FY26 underlying PBT above the base case", float(np.mean(sims["upbt"][0] > base_out["upbt"][0]))),
        ("FY28 underlying PBT above £200m", float(np.mean(sims["upbt"][2] > 200))),
        ("Cash falls below £50m in any year FY26–FY28", float(np.mean(min_cash < 50))),
        ("Cash falls below zero in any year FY26–FY28 (needs new borrowing)", float(np.mean(min_cash < 0))),
    ]
    # tornado: one driver at a time to its min / max (all years), others at base
    tornado = []
    for name in base:
        lo = {k: list(v) for k, v in base.items()}
        hi = {k: list(v) for k, v in base.items()}
        lo[name] = [ranges[name][j][0] for j in range(3)]
        hi[name] = [ranges[name][j][1] for j in range(3)]
        o_lo, o_hi = project(lo, fixed, op), project(hi, fixed, op)
        tornado.append([LABELS[name], round(o_lo["upbt"][2], 1), round(o_hi["upbt"][2], 1),
                        round(o_hi["upbt"][2] - o_lo["upbt"][2], 1),
                        round(min(o_lo["cash"]), 1), round(min(o_hi["cash"]), 1)])
    tornado.sort(key=lambda r: -abs(r[3]))

    with open(outdir / "summary_percentiles.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["measure", "year", "P5", "P25", "P50", "P75", "P95", "mean", "base_case"])
        w.writerows(summary)
    with open(outdir / "probabilities.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["question", "probability"])
        w.writerows([[q, round(p, 3)] for q, p in probs])
    with open(outdir / "tornado_fy28_upbt.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["driver", "FY28 uPBT at driver min", "FY28 uPBT at driver max", "swing",
                    "lowest cash at driver min", "lowest cash at driver max"])
        w.writerows(tornado)
    np.savetxt(outdir / "fy26_upbt_runs.csv", sims["upbt"][0], fmt="%.2f", header="fy26_underlying_pbt", comments="")
    np.savetxt(outdir / "fy28_upbt_runs.csv", sims["upbt"][2], fmt="%.2f", header="fy28_underlying_pbt", comments="")
    np.savetxt(outdir / "min_cash_runs.csv", min_cash, fmt="%.2f", header="min_cash_fy26_fy28", comments="")

    print("\nPercentiles (P5 / P50 / P95):")
    for r in summary:
        print(f"  {r[0]:<30} {r[1]}  {r[2]:>8} {r[4]:>8} {r[6]:>8}   base {r[8]}")
    print("\nProbabilities:")
    for q, p in probs:
        print(f"  {p:6.1%}  {q}")
    print("\nTornado (FY28 underlying PBT):")
    for r in tornado:
        print(f"  {r[0]:<36} {r[1]:>7} to {r[2]:>7}  swing {r[3]}")
    return base_out, sims, summary, probs, tornado, worst, bs_worst


if __name__ == "__main__":
    main()
